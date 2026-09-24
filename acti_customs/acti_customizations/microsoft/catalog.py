# Copyright (c) 2025, Consultoria en Negocios y Aplicaciones and contributors
# For license information, please see license.txt

"""Lectura y validacion del Excel del catalogo Microsoft (precio vigente NCE).

Solo lee/valida el archivo; no toca base de datos. Fail-closed ante archivo mal
formado (hoja o columnas requeridas ausentes): no se hace carga parcial silenciosa.
"""

import openpyxl

SHEET = "Jan_NCE_LicenseBasedPL_GA_MX"

# Columnas requeridas (identidad + comercial + precio). Si falta alguna -> fail-closed.
REQUIRED_COLUMNS = (
	"ProductTitle",
	"ProductId",
	"SkuId",
	"SkuTitle",
	"TermDuration",
	"BillingPlan",
	"Segment",
	"Market",
	"Currency",
	"UnitPrice",
)
# Columnas opcionales que se leen si estan presentes.
OPTIONAL_COLUMNS = ("EffectiveStartDate", "EffectiveEndDate", "Tags", "ChangeIndicator")


class CatalogError(Exception):
	"""El archivo del catalogo no es utilizable (estructura invalida)."""


def _cell(value):
	return "" if value is None else str(value).strip()


def _date(value):
	if value is None or value == "":
		return None
	# openpyxl con data_only devuelve datetime; tambien admite texto.
	text = value.isoformat() if hasattr(value, "isoformat") else str(value)
	return text[:10]


def _float(value):
	if value in (None, ""):
		return None
	try:
		return float(value)
	except TypeError, ValueError:
		return None


def read_catalog(path):
	"""Lee el catalogo y devuelve (rows, header).

	`rows` es una lista de dicts (una por fila de datos) con las claves internas
	product_title, product_id, sku_id, sku_title, term_duration, billing_plan,
	segment, market, currency, unit_price, effective_start_date, effective_end_date,
	tags, change_indicator, y `_row` (numero de fila en el Excel, 1-based).

	Lanza CatalogError si falta la hoja o alguna columna requerida.
	"""
	wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
	if SHEET not in wb.sheetnames:
		wb.close()
		raise CatalogError(f"El archivo no contiene la hoja requerida {SHEET!r}. Hojas: {wb.sheetnames}")
	ws = wb[SHEET]
	all_rows = list(ws.iter_rows(values_only=True))
	wb.close()
	if not all_rows:
		raise CatalogError("La hoja del catalogo esta vacia.")

	header = [str(c).strip() if c is not None else "" for c in all_rows[0]]
	pos = {name: header.index(name) for name in header if name}
	missing = [c for c in REQUIRED_COLUMNS if c not in pos]
	if missing:
		raise CatalogError(f"Faltan columnas requeridas en el catalogo: {missing}")

	def get(row, col):
		i = pos.get(col)
		return row[i] if (i is not None and i < len(row)) else None

	rows = []
	for n, row in enumerate(all_rows[1:], start=2):
		if row is None or all(c is None for c in row):
			continue
		rows.append(
			{
				"_row": n,
				"product_title": _cell(get(row, "ProductTitle")),
				"product_id": _cell(get(row, "ProductId")),
				"sku_id": _cell(get(row, "SkuId")),
				"sku_title": _cell(get(row, "SkuTitle")),
				"term_duration": _cell(get(row, "TermDuration")),
				"billing_plan": _cell(get(row, "BillingPlan")),
				"segment": _cell(get(row, "Segment")),
				"market": _cell(get(row, "Market")),
				"currency": _cell(get(row, "Currency")),
				"unit_price": _float(get(row, "UnitPrice")),
				"effective_start_date": _date(get(row, "EffectiveStartDate")),
				"effective_end_date": _date(get(row, "EffectiveEndDate")),
				"tags": _cell(get(row, "Tags")) or None,
				"change_indicator": _cell(get(row, "ChangeIndicator")) or None,
			}
		)
	return rows, header
